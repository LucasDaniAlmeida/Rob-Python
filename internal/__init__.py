from openpyxl import Workbook, load_workbook
import pandas as pd
import os
import sys
from datetime import datetime

# Variáveis gerais
def _get_writable_base_dir():
    """Return a persistent writable folder for logs.

    In PyInstaller --onefile mode, __file__ points to a temporary _MEI folder,
    so we persist logs next to the executable instead.
    """
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


BASE_DIR = _get_writable_base_dir()
LOG_DIR = os.path.join(BASE_DIR, "log")
LOG_FILE = os.path.join(LOG_DIR, "log.xlsx")
INPUT_FILE = "Cadastrar no SIATE.xlsx"
SHEET_NAME = "Sheet1"

def create_excel_file():
    os.makedirs(LOG_DIR, exist_ok=True)

    if os.path.exists(LOG_FILE):
        try:
            arquivo = load_workbook(LOG_FILE)
            if 'log' not in arquivo.sheetnames:
                planilha = arquivo.create_sheet("log")
                planilha.append(["PE", "FINALIZADO", "DATA_HORA"])
                arquivo.save(LOG_FILE)
            return arquivo
        except Exception as e:
            print(f"Erro ao carregar o arquivo de log: {e}")
            return None
    else:
        print("Arquivo não encontrado, criando um novo.")
        try:
            arquivo = Workbook()
            planilha = arquivo.active
            planilha.title = "log"
            planilha.append(["PE", "FINALIZADO", "DATA_HORA"])
            arquivo.save(LOG_FILE)
            return arquivo
        except Exception as e:
            print(f"Erro ao criar o arquivo de log: {e}")
            return None

def get_log_size():
    try:
        tabela = pd.read_excel(LOG_FILE, sheet_name='log')
        return len(tabela.index) + 1
    except Exception as e:
        print(f"Erro ao obter o tamanho do log: {e}")
        return 0

def write_log(arquivo, Log):
    try:
        # Carregar a planilha de log
        planilha = arquivo['log']
        log_size = planilha.max_row + 1  # Determinar a próxima linha disponível
        data_hora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # Adicionar os dados na próxima linha
        planilha.cell(row=log_size, column=1, value=Log['PE'])
        planilha.cell(row=log_size, column=2, value=Log['FINALIZADO'])
        planilha.cell(row=log_size, column=3, value=data_hora)
        
        # Salvar o arquivo
        arquivo.save(LOG_FILE)
        print("Log salvo com sucesso.")
    except Exception as e:
        print(f"Erro ao salvar o log: {e}")